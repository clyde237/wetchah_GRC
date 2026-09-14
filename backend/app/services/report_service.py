import os
import io
from datetime import datetime
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class GrcReportService:
    @staticmethod
    def generate_audit_mission_pdf(mission, findings, action_plans) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=40,
            rightMargin=40,
            topMargin=40,
            bottomMargin=40
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=6
        )
        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=15
        )
        h2_style = ParagraphStyle(
            'H2Style',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1E293B"),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#334155")
        )
        
        story = []
        
        # En-tête
        story.append(Paragraph(f"RAPPORT DE MISSION D'AUDIT INTERNE — {mission.reference}", title_style))
        story.append(Paragraph(f"Titre : {mission.title} | Statut : {mission.status.upper()} | Date d'édition : {datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0284C7"), spaceAfter=15))
        
        # Périmètre
        story.append(Paragraph("1. Périmètre et Objectifs de la Mission", h2_style))
        story.append(Paragraph(mission.scope or "Aucun périmètre spécifié.", body_style))
        story.append(Spacer(1, 10))
        
        # Constats d'audit (Findings)
        story.append(Paragraph(f"2. Constats d'Audit Relevés ({len(findings)})", h2_style))
        if not findings:
            story.append(Paragraph("Aucun écart ni anomalie relevée au cours de cette mission.", body_style))
        else:
            table_data = [["Réf", "Constat / Libellé", "Sévérité", "Recommandation"]]
            for f in findings:
                table_data.append([
                    f.code,
                    Paragraph(f"<b>{f.title}</b><br/>{f.description[:120]}...", body_style),
                    f.severity.upper(),
                    Paragraph(f.recommendation or "-", body_style)
                ])
            
            t = Table(table_data, colWidths=[60, 200, 75, 180])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ]))
            story.append(t)
            
        story.append(Spacer(1, 15))
        
        # Plans d'action
        story.append(Paragraph(f"3. Plans d'Action Correctifs & Suivi ({len(action_plans)})", h2_style))
        if not action_plans:
            story.append(Paragraph("Aucun plan d'action formel associé pour l'instant.", body_style))
        else:
            plan_data = [["Action", "Échéance", "Statut"]]
            for p in action_plans:
                plan_data.append([
                    Paragraph(f"<b>{p.title}</b><br/>{p.description}", body_style),
                    p.due_date.strftime('%d/%m/%Y') if p.due_date else "-",
                    p.status.upper()
                ])
            pt = Table(plan_data, colWidths=[330, 95, 90])
            pt.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            story.append(pt)

        doc.build(story)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_risks_excel(risks) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cartographie Risques"
        
        headers = [
            "Code", "Titre", "Catégorie", "Processus", 
            "Impact Brut", "Vraisemblance Brute", "Score Brut",
            "Impact Résiduel", "Vraisemblance Résiduelle", "Score Résiduel",
            "Stratégie", "Statut"
        ]
        ws.append(headers)
        
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        for r in risks:
            ws.append([
                r.code, r.title, r.category, r.process_affected or "-",
                r.gross_impact, r.gross_likelihood, r.gross_score,
                r.residual_impact, r.residual_likelihood, r.residual_score,
                r.treatment_strategy, r.status
            ])
            
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
